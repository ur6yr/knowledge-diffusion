"""Versioned discovery cohorts and a bounded durable expansion frontier."""

import csv
import hashlib
from pathlib import Path

from kdiff.core.contracts import canonical, digest
from kdiff.core.durable import TaskLedger

VENUE_PROFILES = {'paper-2025': (2020, 2024), 'qualifying-2026': (2010, 2024)}


def import_csv(path, store, ledger, *, release, name_column, id_column=None,
               max_records=1000, max_bytes=16*1024*1024):
    path=Path(path)
    if path.is_symlink() or not path.is_file() or path.stat().st_size > max_bytes:
        raise ValueError('Seed file unavailable or exceeds budget')
    if not release:
        raise ValueError('Seed release label is required')
    capture=store.capture(path.read_bytes())
    seeds=[]
    with path.open(encoding='utf-8-sig',newline='') as stream:
        rows=csv.DictReader(stream)
        if name_column not in (rows.fieldnames or []) or id_column and id_column not in rows.fieldnames:
            raise ValueError('Seed columns do not match the inspected source header')
        for line,row in enumerate(rows,2):
            if len(seeds)>=max_records:
                raise ValueError('Seed record budget exceeded, no partial successful cohort')
            if not row[name_column].strip():
                raise ValueError('Seed has no name')
            seed={'kind':'researcher','name':row[name_column].strip(),
                  'source_id':row.get(id_column) if id_column else None,'release':release,
                  'source_artifact':capture['sha256'],'source_line':line,'identity_verified':False,
                  'depth':0,'included_by':'discovery_cohort'}
            seeds.append(seed)
    if not seeds:
        raise ValueError('Seed cohort empty')
    ids=[ledger.enqueue(seed) for seed in seeds]
    return {'release':release,'records':len(ids),'task_ids':ids,'raw_artifact':capture['sha256'],
            'scope':'Supplied discovery cohort only, not verified person identities'}


def enqueue_expansion(ledger, parent, candidates, *, max_depth, max_entities, relations,
                      years, profile=None):
    if not 0<=max_depth<=5 or not 1<=max_entities<=10000:
        raise ValueError('Expansion budgets invalid')
    if profile and tuple(years)!=VENUE_PROFILES.get(profile):
        raise ValueError('Venue seed window conflicts with named historical profile')
    state,revision=ledger.read_checkpoint('frontier')
    state=state or {'seen':[],'decisions':[]}
    additions=[]
    for row in candidates:
        key=digest([row['kind'],row['source_id'],row.get('source_version')])
        reason='included'
        if parent.get('depth',0)>=max_depth:
            reason='depth_limit'
        elif row.get('relation') not in relations:
            reason='relation_scope'
        elif row.get('year') is None or not years[0]<=row['year']<=years[1]:
            reason='unknown_or_outside_year_scope'
        elif key in state['seen']:
            reason='already_seen'
        elif len(state['seen'])>=max_entities:
            reason='entity_limit'
        if reason=='included':
            if not row.get('source_artifact'):
                raise ValueError('Expansion inclusion needs source provenance')
            child={**row,'depth':parent.get('depth',0)+1,'included_by':digest(parent)}
            # Task identity includes the input content and parent derivation.
            additions.append(ledger.enqueue(child))
            state['seen'].append(key)
        state['decisions'].append({'candidate':key,'reason':reason})
    ledger.checkpoint('frontier',state,revision)
    return {'enqueued':additions,'decisions':state['decisions'][-len(candidates):] if candidates else []}


def import_xlsx(path, store, ledger, *, release, sheet, name_column, id_column=None,
                limit=1000, max_file_bytes=128*1024*1024, max_uncompressed_bytes=512*1024*1024):
    """Import an explicitly scoped prefix of an existing workbook as discovery seeds.

    The source file is hashed in bounded chunks. Selected cells are captured as
    immutable evidence with sheet/row addresses. No citation indicators or names
    are treated as verified person identifiers.
    """
    from zipfile import ZipFile
    from openpyxl import load_workbook
    path = Path(path)
    if path.is_symlink() or not path.is_file() or path.stat().st_size > max_file_bytes:
        raise ValueError('Seed workbook exceeds the declared file budget or is unavailable')
    if not release or not 1 <= limit <= 1000:
        raise ValueError('Workbook release and bounded prefix length are required')
    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > max_uncompressed_bytes:
            raise ValueError('Workbook expansion exceeds the declared byte budget')
    checksum = hashlib.sha256()
    before = path.stat()
    with path.open('rb') as stream:
        while block := stream.read(1024*1024):
            checksum.update(block)
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if sheet not in book.sheetnames:
            raise ValueError('Requested worksheet is absent')
        rows = book[sheet].iter_rows(values_only=True)
        header = list(next(rows))
        if header.count(name_column) != 1 or id_column and header.count(id_column) != 1:
            raise ValueError('Workbook columns do not match the inspected header')
        seeds = []
        for number in range(2, limit + 2):
            values = next(rows, None)
            if values is None:
                break
            raw = dict(zip(header, values))
            name = raw[name_column]
            if not isinstance(name, str) or not name.strip():
                raise ValueError('Selected workbook row has no author name')
            captured = store.capture(canonical({'source_file_sha256': checksum.hexdigest(), 'sheet': sheet,
                                                'row': number, 'cells': raw}))
            seeds.append({'kind': 'researcher', 'name': name.strip(), 'source_id': raw.get(id_column) if id_column else None,
                          'release': release, 'source_artifact': captured['sha256'], 'source_line': number,
                          'source_sheet': sheet, 'source_file_sha256': checksum.hexdigest(),
                          'identity_verified': False, 'depth': 0, 'included_by': 'discovery_cohort'})
    finally:
        book.close()
    after = path.stat()
    if (before.st_mtime_ns, before.st_size) != (after.st_mtime_ns, after.st_size):
        raise ValueError('Workbook changed during capture')
    task_ids = [ledger.enqueue(seed) for seed in seeds]
    report = {'release': release, 'source_file_sha256': checksum.hexdigest(), 'source_file_bytes': before.st_size,
              'sheet': sheet, 'requested_prefix_rows': limit, 'records': len(seeds), 'task_ids': task_ids,
              'scope': 'Only the explicitly selected worksheet prefix. Discovery cohort, not verified identities.'}
    return {**report, 'manifest_id': store.put(report)}
